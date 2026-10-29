import sys, os, json, glob, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule

src_dir, out = sys.argv[1], sys.argv[2]
today = dt.date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else dt.date.today()
docs = []
for p in sorted(glob.glob(os.path.join(src_dir, "*.json"))):
    d = json.load(open(p, encoding="utf-8"))
    d = d.get("data", d) if isinstance(d.get("data"), dict) else d
    d["_id"] = os.path.basename(p)[:-5]
    docs.append(d)
docs.sort(key=lambda d: (d.get("created", ""), d["_id"]))
TEMP = {"hot": "🔥 Горячий", "warm": "🌤 Тёплый", "cold": "❄ Холодный", "dealer": "✅ Дилер", "branch": "↪ Филиал", "refused": "⛔ Отказ"}
def dd(s):
    try: return dt.date.fromisoformat(s) if s else None
    except Exception: return s
npairs = max([6] + [len([t for t in d.get("touches", []) if t and t.get("t")]) for d in docs])
head = ["№", "Компания", "ИНН", "Выручка 2025, млн ₽", "ФИО", "Должность", "Регион", "Телефон", "Email", "Статус", "канал связи", "знакомство/намерения"]
for i in range(1, npairs + 1): head += [f"дата {i} касания", f"Итоги {i} касания"]
head += ["следующая дата созвона", "Статус клиента (WAYTEKO)", "Следующий шаг (план)", "Встречи (предстоящие)"]
NEXT = len(head) - 3
widths = [5, 34, 13, 13, 26, 45, 24, 34, 34, 29, 17, 55] + [15, 45] * npairs + [23, 20, 45, 40]

wb = Workbook(); ws = wb.active; ws.title = "Визитки"
thin = Side(style="thin", color="000000"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
for c, (h, w) in enumerate(zip(head, widths), 1):
    x = ws.cell(1, c, h); x.border = B
    x.font = Font(name="Calibri", size=14, bold=True, color="000000" if c == NEXT else "FFFFFF")
    x.fill = PatternFill("solid", fgColor="FFFF00" if c == NEXT else "1F4E78")
    x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions[L(c)].width = w
ws.row_dimensions[1].height = 76.5
for r, d in enumerate(docs, 2):
    touches = [t for t in d.get("touches", []) if t and t.get("t")]
    row = [r - 1, d.get("name"), d.get("inn"), d.get("revenue"), d.get("fio"), d.get("position"), d.get("region"),
           d.get("phone"), d.get("email"), d.get("status"), d.get("channel"), d.get("intro")]
    for i in range(npairs):
        t = touches[i] if i < len(touches) else {}
        row += [dd(t.get("d")), t.get("t")]
    meets = sorted([m for m in d.get("meets", []) or [] if m and m.get("d") and not m.get("done") and m["d"] >= today.isoformat()], key=lambda m: (m["d"], m.get("t") or ""))
    mtxt = "\n".join(f'{dt.date.fromisoformat(m["d"]).strftime("%d.%m")}{" " + m["t"] if m.get("t") else ""}{" — " + m["place"] if m.get("place") else ""}{": " + m["note"] if m.get("note") else ""}' for m in meets)
    row += [dd(d.get("next")), TEMP.get(d.get("temp"), d.get("temp")), d.get("plan"), mtxt or None]
    temp = d.get("temp")
    for c, v in enumerate(row, 1):
        x = ws.cell(r, c, v if v != "" else None); x.border = B
        x.font = Font(name="Calibri", size=14, color="A6A6A6" if temp in ("branch", "refused") else "000000", bold=(c == NEXT + 1))
        x.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(v, dt.date): x.number_format = "d-mmm"
        if temp == "hot": x.fill = PatternFill("solid", fgColor="00B0F0")
        elif temp == "dealer": x.fill = PatternFill("solid", fgColor="A9D08E")
        if c == NEXT and temp not in ("hot", "dealer"): x.fill = PatternFill("solid", fgColor="FFFF00")
    # последний итог — зелёным
    if touches:
        c = 12 + 2 * len(touches)
        ws.cell(r, c).fill = PatternFill("solid", fgColor="C6EFCE"); ws.cell(r, c).font = Font(name="Calibri", size=14, bold=True, color="006100")
last = len(docs) + 1
nl = L(NEXT)
rng = f"A2:{L(len(head))}{max(last, 2)}"
act = f'AND($A2<>"",ISNUMBER(${nl}2),${L(NEXT + 1)}2<>"↪ Филиал",${L(NEXT + 1)}2<>"⛔ Отказ")'
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({act},${nl}2<TODAY())'], fill=PatternFill("solid", fgColor="FF7C80"), stopIfTrue=True))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({act},${nl}2=TODAY())'], fill=PatternFill("solid", fgColor="FFE699"), stopIfTrue=True))
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{L(len(head))}{last}"
wb.save(out)
print(f"OK: {len(docs)} клиентов, {npairs} пар касаний -> {out}")
