import json, sys
from playwright.sync_api import sync_playwright
PW = open(sys.argv[1]).read().strip()
errs = []
with sync_playwright() as p:
    b = p.chromium.launch(); ctx = b.new_context(viewport={"width":390,"height":844}, accept_downloads=True)
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto("http://localhost:8765/"); pg.wait_for_timeout(1500)
    print("1 экран входа:", pg.locator("#pwForm").count())
    pg.fill("#pw", "неверный"); pg.click("#pwForm button[type=submit]"); pg.wait_for_timeout(2500)
    print("2 после неверного пароля форма осталась:", pg.locator("#pwForm").count(), "| toast:", pg.locator("#toast").inner_text())
    pg.fill("#pw", PW); pg.click("#pwForm button[type=submit]"); pg.wait_for_timeout(3000)
    n = len(json.loads(pg.evaluate("localStorage.getItem('wayteko.clients.v1')")))
    print("3 загружено клиентов:", n, "| шапка:", pg.locator("#todayLbl").inner_text())
    pg.screenshot(path=sys.argv[2]+"/s1.png")
    pg.reload(); pg.wait_for_timeout(3000)
    print("4 после перезагрузки без пароля, форма входа:", pg.locator("#pwForm").count(), "клиентов на экране:", pg.locator(".card").count())
    pg.click("#dataBtn"); pg.wait_for_timeout(300)
    pg.screenshot(path=sys.argv[2]+"/s2.png")
    with pg.expect_download(timeout=15000) as d:
        pg.click("button[data-act=xlsx]")
    dl = d.value; path = sys.argv[2]+"/dl.xlsx"; dl.save_as(path)
    print("5 скачан файл:", dl.suggested_filename)
    b.close()
import openpyxl
ws = openpyxl.load_workbook(path).active
print("6 строк в Excel:", ws.max_row-1, "| последняя колонка:", ws.cell(1, ws.max_column).value)
print("ERRORS:", errs or "нет")
